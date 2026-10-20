"""Read-only XLSX verification. Does not author or modify workbooks."""
import json, pathlib, math
import openpyxl
root=pathlib.Path(__file__).resolve().parents[1]
for name in ['baseline-budget-request','changed-budget-request']:
 path=root/'docs/verification'/f'{name}.xlsx'
 expected=json.loads(path.with_suffix('.json').read_text())
 values=openpyxl.load_workbook(path,data_only=True)
 formulas=openpyxl.load_workbook(path,data_only=False)
 assert values.sheetnames==['예산 집행 승인 요청서','노무비 산출 상세','시나리오 비교','계산 근거 및 가정']
 req=values.worksheets[0]; data=expected['result']
 assert req['B13'].value==data['incremental']
 assert req['B14'].value==data['total']
 assert req['B15'].value==data['budgetRemaining']
 assert all(req[f'B{i}'].value in [None,''] for i in range(18,23))
 detail=values.worksheets[1]
 assert detail.cell(detail.max_row,12).value==data['total']
 assert sum(detail.cell(i,12).value or 0 for i in range(2,detail.max_row-2))==data['incremental']
 for i,row in enumerate(expected['result']['costs'],2):
  assert math.isclose(detail.cell(i,7).value,row['headcount']*row['hours'],abs_tol=1e-7)
  assert formulas.worksheets[1].cell(i,7).value==f'=E{i}*F{i}'
 assert formulas.worksheets[1].cell(detail.max_row,12).data_type=='f'
 assert all(s.freeze_panes=='A2' for s in values.worksheets)
 print(name,'PASS: four sheets, cached formulas, costs, budget, blank approvals, frozen headers')
